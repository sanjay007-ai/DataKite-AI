"""Regression tests for DataKite's universal file reader."""
from pathlib import Path
import pandas as pd
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from backend.file_reader import read_file


def make_fixture(folder: Path):
    df = pd.DataFrame({
        "Order_ID": [1, 2, 3],
        "Order_Date": ["2026-01-01", "2026-02-01", "2026-03-01"],
        "Product": ["A", "B", "A"],
        "Sales": [100, 200, 150],
        "Profit": [20, 30, 25],
    })
    df.to_csv(folder / "data.csv", index=False)
    with pd.ExcelWriter(folder / "data.xlsx") as writer:
        pd.DataFrame({"title": ["cover"]}).to_excel(writer, index=False, sheet_name="Cover")
        df.to_excel(writer, index=False, sheet_name="Business Data")
    doc = Document()
    table = doc.add_table(rows=1, cols=len(df.columns))
    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = col
    for row in df.itertuples(index=False):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    doc.save(folder / "data.docx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_table(len(df) + 1, len(df.columns), Inches(1), Inches(1), Inches(8), Inches(3))
    for i, col in enumerate(df.columns):
        shape.table.cell(0, i).text = col
    for r, row in enumerate(df.itertuples(index=False), 1):
        for c, value in enumerate(row):
            shape.table.cell(r, c).text = str(value)
    prs.save(folder / "data.pptx")


def test_common_formats(tmp_path):
    make_fixture(tmp_path)
    for name in ("data.csv", "data.xlsx", "data.docx", "data.pptx"):
        result = read_file(str(tmp_path / name))
        assert result.get("data") is not None, name
        assert len(result["data"]) == 3, name
        assert len(result["data"].columns) == 5, name


def test_messy_excel_and_csv_header_detection(tmp_path):
    rows = [
        ["TECHNOVA RAW DATA"],
        ["Generated report"],
        ["Order ID", "Order Date", "Product", "Sales", "Profit"],
        [1, "2026-01-01", "A", 100, 20],
        [2, "2026-01-02", "B", 200, 30],
    ]
    raw = pd.DataFrame(rows)
    raw.to_csv(tmp_path / "messy.csv", index=False, header=False)
    with pd.ExcelWriter(tmp_path / "messy.xlsx") as writer:
        raw.to_excel(writer, index=False, header=False, sheet_name="Raw")
    for name in ("messy.csv", "messy.xlsx"):
        result = read_file(str(tmp_path / name))
        assert list(result["data"].columns)[:5] == ["Order ID", "Order Date", "Product", "Sales", "Profit"]
        assert len(result["data"]) == 2
