import pandas as pd
from pathlib import Path
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "backend"))
sys.path.insert(0, str(PROJECT_ROOT))
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from upload_pipeline import parse_uploaded_file


def make_business_df():
    return pd.DataFrame({
        "Order ID": ["O1", "O2", "O3"],
        "Order Date": ["2026-01-01", "2026-01-02", "2026-01-03"],
        "Customer": ["C1", "C2", "C1"],
        "City": ["Chennai", "Mumbai", "Chennai"],
        "Sales": [1000, 2000, 1500],
        "Profit": [100, 300, 200],
    })


def test_exact_http_pipeline_for_common_formats(tmp_path):
    df = make_business_df()
    files = []
    csv = tmp_path / "sales.csv"; df.to_csv(csv, index=False); files.append(csv)
    xlsx = tmp_path / "sales.xlsx"; df.to_excel(xlsx, index=False); files.append(xlsx)
    doc = Document(); t = doc.add_table(rows=1, cols=len(df.columns))
    for i,c in enumerate(df.columns): t.rows[0].cells[i].text=c
    for row in df.itertuples(index=False):
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    docx=tmp_path/'sales.docx'; doc.save(docx); files.append(docx)
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[5]); table=slide.shapes.add_table(len(df)+1,len(df.columns),Inches(.2),Inches(.2),Inches(10),Inches(3)).table
    for i,c in enumerate(df.columns): table.cell(0,i).text=c
    for r,row in enumerate(df.itertuples(index=False),1):
        for i,v in enumerate(row): table.cell(r,i).text=str(v)
    pptx=tmp_path/'sales.pptx'; prs.save(pptx); files.append(pptx)
    for path in files:
        result=parse_uploaded_file(str(path),path.name)
        assert result['accepted'], (path.name,result)
        assert result['data'].shape[0] == 3
        assert 'Sales' in result['data'].columns


def test_exact_http_pipeline_messy_csv(tmp_path):
    path=tmp_path/'technova_raw.csv'
    path.write_text('TECHNOVA RAW DATA\nGenerated report\n\nOrder ID,Order Date,Customer,Sales\nO1,2026-01-01,C1,100\nO2,2026-01-02,C2,200\n',encoding='utf-8')
    result=parse_uploaded_file(str(path),path.name)
    assert result['accepted']
    assert list(result['data'].columns)[:4] == ['Order ID','Order Date','Customer','Sales']
    assert len(result['data']) == 2
