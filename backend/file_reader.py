"""Universal, fault-tolerant business file reader for DataKite AI.

The reader's contract is deliberately simple: structured files return a dict
containing a pandas DataFrame under ``data``.  Text-only documents return
``text`` as a fallback.  The caller never needs to know which parser was used.
"""
import os
import json
import re
import pandas as pd


def _unique_columns(columns):
    out, seen = [], {}
    for i, col in enumerate(columns):
        name = str(col).strip() if col is not None else ""
        if not name or name.lower().startswith("unnamed:"):
            name = f"Column_{i + 1}"
        key = name.casefold()
        seen[key] = seen.get(key, 0) + 1
        out.append(name if seen[key] == 1 else f"{name}_{seen[key]}")
    return out


def clean_columns(df):
    if not isinstance(df, pd.DataFrame):
        return None
    df = df.copy()
    df = df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    if df.empty:
        return None
    df.columns = _unique_columns(df.columns)
    return df


def _looks_like_header(row):
    vals = [str(x).strip() for x in row if pd.notna(x) and str(x).strip()]
    if len(vals) < 2:
        return False
    unique = {v.casefold() for v in vals}
    if len(unique) < max(2, int(len(vals) * 0.6)):
        return False
    business_words = (
        "id", "date", "sales", "sale", "revenue", "amount", "profit", "cost",
        "product", "customer", "category", "quantity", "qty", "city", "region",
        "status", "payment", "order", "invoice", "department", "price", "discount",
        "time", "state", "country", "phone", "email"
    )
    text = " ".join(vals).casefold()
    semantic = sum(1 for word in business_words if word in text)
    # A row made mostly of numbers is data, not a header.
    numeric_like = sum(bool(re.fullmatch(r"[-+]?\d+(?:\.\d+)?", v.replace(",", ""))) for v in vals)
    return semantic > 0 or numeric_like < max(1, len(vals) // 2)


def _table_from_header(raw, header_index):
    body = raw.iloc[header_index + 1:].copy()
    header = ["" if pd.isna(x) else str(x).strip() for x in raw.iloc[header_index].tolist()]
    if len(header) != len(body.columns):
        return None
    body.columns = _unique_columns(header)
    body = body.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    if body.empty or len(body.columns) < 1:
        return None
    # Drop rows that are completely blank strings after parsing.
    body = body.apply(lambda col: col.map(lambda v: pd.NA if (isinstance(v, str) and not v.strip()) else v))
    body = body.dropna(how="all").reset_index(drop=True)
    if body.empty:
        return None
    return body


def _usable_table(df, min_rows=1, min_cols=1):
    if not isinstance(df, pd.DataFrame):
        return None
    df = clean_columns(df)
    if df is None or len(df) < min_rows or len(df.columns) < min_cols:
        return None
    # At least one meaningful header and one non-empty data cell.
    if not any(str(c).strip() and not str(c).startswith("Column_") for c in df.columns):
        return None
    if not df.notna().any().any():
        return None
    return df


def _infer_table_from_raw(raw, max_header_scan=30):
    """Find the most likely header row in a raw matrix.

    This handles CSV/Excel exports with title rows, blank rows, merged-title
    artifacts, report headers, and numeric column labels.
    """
    if not isinstance(raw, pd.DataFrame):
        return None
    raw = raw.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    if raw.empty:
        return None

    candidates = []
    limit = min(max_header_scan, max(0, len(raw) - 2))
    for i in range(limit + 1):
        row = raw.iloc[i].tolist()
        if not _looks_like_header(row):
            continue
        body = _table_from_header(raw, i)
        if body is None or len(body) < 1:
            continue
        header_text = " ".join(str(x).strip().casefold() for x in row if pd.notna(x))
        business_words = ("id", "date", "sales", "revenue", "amount", "profit", "product", "customer", "category", "quantity", "city", "region", "status", "payment", "order")
        semantic = sum(8 for w in business_words if w in header_text)
        nonempty = sum(bool(str(x).strip()) for x in row if pd.notna(x))
        score = semantic + nonempty * 10 + min(len(body), 200)
        # Prefer a header that leaves a substantial body over a later accidental row.
        if i == 0:
            score += 20
        candidates.append((score, body))
    if candidates:
        return max(candidates, key=lambda x: x[0])[1]

    # Last-resort: if a normal parser already produced useful headers, keep it.
    direct = _usable_table(raw)
    if direct is not None:
        cols = [str(c).casefold() for c in direct.columns]
        if not all(c.isdigit() or c.startswith("column_") for c in cols):
            return direct
    return None


def _rows_to_dataframe(rows):
    cleaned = []
    for row in rows or []:
        if row is None:
            continue
        vals = ["" if v is None else str(v).strip() for v in row]
        if any(vals):
            cleaned.append(vals)
    if len(cleaned) < 2:
        return None
    width = max(len(r) for r in cleaned)
    raw = pd.DataFrame([r + [""] * (width - len(r)) for r in cleaned])
    return _infer_table_from_raw(raw)


# CSV -----------------------------------------------------------------------
def _detect_delimiter(text):
    lines = [ln for ln in text.splitlines() if ln.strip()][:120]
    best = None
    for delimiter in [",", "\t", ";", "|"]:
        counts = [ln.count(delimiter) for ln in lines]
        positive = [c for c in counts if c > 0]
        if not positive:
            continue
        # Data rows usually have the same delimiter count. Ignore title rows
        # that contain no delimiter instead of letting them poison detection.
        max_count = max(positive)
        consistency = sum(1 for c in counts if c == max_count)
        score = consistency * 100 + max_count
        if best is None or score > best[0]:
            best = (score, delimiter)
    return best[1] if best else ","


def _read_csv_raw(file_path, encoding):
    import csv
    text = open(file_path, "r", encoding=encoding, errors="replace", newline="").read()
    delimiter = _detect_delimiter(text)
    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    rows = [r for r in rows if any(str(v).strip() for v in r)]
    if not rows:
        return pd.DataFrame()
    width = max(len(r) for r in rows)
    return pd.DataFrame([r + [""] * (width - len(r)) for r in rows])


def read_csv(file_path):
    # Fast path for normal business CSVs. The previous implementation always
    # loaded the entire file with csv.reader and then rebuilt multiple
    # DataFrames for header inference. That is unnecessarily expensive on
    # cloud/serverless CPUs. Keep the smart fallback for messy report exports.
    errors = []
    for enc in ["utf-8-sig", "utf-8", "cp1252", "latin1"]:
        try:
            direct = pd.read_csv(file_path, encoding=enc)
            direct = _usable_table(direct)
            if direct is not None:
                cols = [str(c).strip().casefold() for c in direct.columns]
                meaningful = [c for c in cols if not c.isdigit() and not c.startswith("unnamed") and not c.startswith("column_") and c]
                # A single named cell followed by Unnamed columns is usually a
                # report title, not the real header. Let the smart fallback find
                # the actual header row in that case.
                if len(meaningful) >= 2:
                    return direct
        except Exception as exc:
            errors.append(str(exc))
        try:
            raw = _read_csv_raw(file_path, enc)
            usable = _infer_table_from_raw(raw)
            if usable is not None:
                return usable
        except Exception as exc:
            errors.append(str(exc))
    raise ValueError("CSV could not be read as a usable table. The file may be empty, malformed, or missing a header/data row.")


# Excel ---------------------------------------------------------------------
def read_excel(file_path):
    try:
        book = pd.ExcelFile(file_path)
    except Exception as exc:
        raise ValueError(f"Excel file could not be opened: {exc}") from exc
    candidates = []
    for sheet in book.sheet_names:
        try:
            # Fast path: normal Excel business tables are by far the common
            # case. Avoid reading every sheet as a raw matrix and scanning
            # every cell unless the normal header parser fails.
            direct = pd.read_excel(book, sheet_name=sheet)
            usable_direct = _usable_table(direct)
            if usable_direct is not None:
                cols = [str(c).strip().casefold() for c in usable_direct.columns]
                meaningful = [c for c in cols if not c.isdigit() and not c.startswith("unnamed") and not c.startswith("column_") and c]
                if len(meaningful) >= 2:
                    usable_direct.attrs["source_sheet"] = sheet
                    header_text = " ".join(cols)
                    semantic = sum(8 for w in ("id", "date", "sales", "revenue", "amount", "profit", "product", "customer", "category", "quantity", "city", "region", "status", "payment", "order") if w in header_text)
                    candidates.append((semantic + len(usable_direct) * max(1, len(usable_direct.columns)), sheet, usable_direct))
                    continue

            # Fallback only for messy/report-style sheets with title rows or
            # unusual headers.
            raw = pd.read_excel(book, sheet_name=sheet, header=None)
            usable = _infer_table_from_raw(raw)
            if usable is not None:
                usable.attrs["source_sheet"] = sheet
                header_text = " ".join(str(c).casefold() for c in usable.columns)
                semantic = sum(8 for w in ("id", "date", "sales", "revenue", "amount", "profit", "product", "customer", "category", "quantity", "city", "region", "status", "payment", "order") if w in header_text)
                candidates.append((semantic + len(usable) * max(1, len(usable.columns)), sheet, usable))
        except Exception as exc:
            print(f"[DataKite] Excel sheet skipped: {sheet}: {exc}")
    if candidates:
        _, sheet, df = max(candidates, key=lambda x: x[0])
        df.attrs["source_sheet"] = sheet
        return df
    raise ValueError("Excel file contains no usable table. Check that at least one sheet has column headers and data rows.")


# JSON ----------------------------------------------------------------------
def read_json(file_path):
    with open(file_path, "r", encoding="utf-8-sig") as file:
        obj = json.load(file)
    if isinstance(obj, list):
        df = pd.json_normalize(obj) if obj and isinstance(obj[0], dict) else pd.DataFrame(obj)
    elif isinstance(obj, dict):
        records = next((v for v in obj.values() if isinstance(v, list) and v and isinstance(v[0], dict)), None)
        df = pd.json_normalize(records) if records is not None else pd.DataFrame([obj])
    else:
        raise ValueError("JSON does not contain tabular data.")
    df = _usable_table(df)
    if df is None:
        raise ValueError("JSON contains no usable tabular data.")
    return df


# PDF -----------------------------------------------------------------------
def read_pdf(file_path, max_pages=200):
    import pdfplumber
    tables, text_parts = [], []
    with pdfplumber.open(file_path) as pdf:
        total_pages = len(pdf.pages)
        for page in pdf.pages[:max_pages]:
            text = page.extract_text() or ""
            if text:
                text_parts.append(text)
            for table in page.extract_tables() or []:
                df = _rows_to_dataframe(table)
                if df is not None:
                    tables.append(df)
    if tables:
        data = max(tables, key=lambda d: len(d) * max(1, len(d.columns)))
        return {"type": "table", "data": data, "text": "\n".join(text_parts), "truncated": total_pages > max_pages, "pages_read": min(total_pages, max_pages), "total_pages": total_pages}
    lines = [ln.strip() for text in text_parts for ln in text.splitlines() if ln.strip()]
    rows = []
    for ln in lines:
        parts = [p.strip() for p in re.split(r"\s{2,}|\t|\|", ln) if p.strip()]
        if len(parts) >= 2:
            rows.append(parts)
    data = _rows_to_dataframe(rows)
    if data is not None and len(data.columns) >= 2:
        return {"type": "table", "data": data, "text": "\n".join(text_parts), "truncated": total_pages > max_pages, "pages_read": min(total_pages, max_pages), "total_pages": total_pages}
    return {"type": "text", "data": None, "text": "\n".join(text_parts), "truncated": total_pages > max_pages, "pages_read": min(total_pages, max_pages), "total_pages": total_pages}


# DOCX ----------------------------------------------------------------------
def read_word(file_path):
    from docx import Document
    document = Document(file_path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    tables = []
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        df = _rows_to_dataframe(rows)
        if df is not None:
            tables.append(df)
    # Also recognize tab-delimited / pipe-delimited text pasted into Word.
    if not tables and paragraphs:
        rows = []
        for line in paragraphs:
            parts = [p.strip() for p in re.split(r"\t|\s{2,}|\|", line) if p.strip()]
            if len(parts) >= 2:
                rows.append(parts)
        df = _rows_to_dataframe(rows)
        if df is not None:
            tables.append(df)
    data = max(tables, key=lambda d: len(d) * max(1, len(d.columns))) if tables else None
    return {"type": "table" if data is not None else "document", "data": data, "text": "\n".join(paragraphs), "tables": [t.to_dict(orient="records") for t in tables]}


# PPTX ----------------------------------------------------------------------
def read_pptx(file_path):
    from pptx import Presentation
    presentation = Presentation(file_path)
    slides, text_parts, tables = [], [], []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text = shape.text.strip()
                slide_text.append(text)
                text_parts.append(text)
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                df = _rows_to_dataframe(rows)
                if df is not None:
                    tables.append(df)
        slides.append({"slide": slide_number, "text": "\n".join(slide_text)})
    data = max(tables, key=lambda d: len(d) * max(1, len(d.columns))) if tables else None
    return {"type": "table" if data is not None else "presentation", "data": data, "text": "\n".join(text_parts), "slides": slides}


def read_file(file_path):
    if not file_path or not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".csv":
        return {"type": "dataset", "format": "csv", "data": read_csv(file_path)}
    if extension in {".xlsx", ".xls"}:
        return {"type": "dataset", "format": "excel", "data": read_excel(file_path)}
    if extension == ".json":
        return {"type": "dataset", "format": "json", "data": read_json(file_path)}
    if extension == ".pdf":
        return read_pdf(file_path)
    if extension == ".docx":
        return read_word(file_path)
    if extension == ".pptx":
        return read_pptx(file_path)
    raise ValueError(f"Unsupported file type: {extension}")


def get_file_info(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError("File not found.")
    return {"filename": os.path.basename(file_path), "extension": os.path.splitext(file_path)[1].lower(), "size_bytes": os.path.getsize(file_path)}
